"""
Render business-case content blocks to an A4 Word document with python-docx.
Square-bracketed text, e.g. [Trust name], is highlighted yellow for the trust
to complete. No horizontal rules; headings use built-in styles so the
contents field works. Every page carries a draft header until the
assumptions have been reviewed.
"""
import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
DRAFT_RED = RGBColor(0xC0, 0x00, 0x00)
HEADER_TEXT = "DRAFT: assumptions awaiting specialty review. Not for approval."
GREY = RGBColor(0x59, 0x59, 0x59)
STATUS_FILL = {"Green": "C6EFCE", "Amber": "FFEB9C", "Red": "FFC7CE"}
FONT = "Arial"
CONTENT_DXA = 9026


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "BFBFBF")
        b.append(e)
    tblPr.append(b)


def _runs(par, text, size=11, bold=False, color=None):
    for part in [p for p in re.split(r"(\[[^\]]+\])", str(text)) if p]:
        r = par.add_run(part)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
        if color is not None:
            r.font.color.rgb = color
        if part.startswith("["):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _field(par, instr):
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        r = par.add_run()
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def _table(doc, b):
    widths = b.get("widths") or [CONTENT_DXA // len(b["header"])] * len(b["header"])
    size = 8.5 if b.get("small") else 9.5
    t = doc.add_table(rows=1, cols=len(b["header"]))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _borders(t)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for i, h in enumerate(b["header"]):
        c = t.rows[0].cells[i]
        _shade(c, "1F3864")
        c.paragraphs[0].text = ""
        _runs(c.paragraphs[0], h, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for ri, row in enumerate(b["rows"]):
        cells = t.add_row().cells
        last = b.get("bold_last") and ri == len(b["rows"]) - 1
        for i, v in enumerate(row):
            if b.get("status_col") == i and v in STATUS_FILL:
                _shade(cells[i], STATUS_FILL[v])
            elif ri % 2:
                _shade(cells[i], "F2F2F2")
            _runs(cells[i].paragraphs[0], v, size=size, bold=bool(last))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w / 567)
            for par in row.cells[i].paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
    doc.add_paragraph()


def render(blocks):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    st = doc.styles
    st["Normal"].font.name = FONT
    st["Normal"].font.size = Pt(11)
    for name, size, col in (("Heading 1", 16, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11, None)):
        s = st[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = col if col is not None else RGBColor(0, 0, 0)
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), FONT)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _runs(hp, HEADER_TEXT, size=8, color=DRAFT_RED)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(fp, "Page ", size=8, color=GREY)
    _field(fp, "PAGE")

    for b in blocks:
        t = b["type"]
        if t == "cover":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(110)
            _runs(p, b["trust"], size=16, bold=True, color=NAVY)
            _runs(doc.add_paragraph(), b["title"], size=26, bold=True)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(40)
            _runs(p, b["subtitle"], size=14, color=GREY)
            ct = doc.add_table(rows=0, cols=2)
            _borders(ct)
            for k, v in b["fields"]:
                cells = ct.add_row().cells
                _shade(cells[0], "F2F2F2")
                _runs(cells[0].paragraphs[0], k, bold=True)
                _runs(cells[1].paragraphs[0], v)
                cells[0].width, cells[1].width = Cm(5.3), Cm(10.6)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(40)
            _runs(p, b["footnote"], size=9, color=GREY)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t == "toc":
            doc.add_heading("Contents", level=1)
            _field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
            _runs(doc.add_paragraph(), "[Right-click and choose Update Field to show the contents.]", size=9)
        elif t == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif t in ("h1", "h2", "h3"):
            doc.add_heading(b["text"], level=int(t[1]))
        elif t == "p":
            _runs(doc.add_paragraph(), b["text"])
        elif t == "bullets":
            for item in b["items"]:
                _runs(doc.add_paragraph(style="List Bullet"), item)
        elif t == "numbered":
            for item in b["items"]:
                _runs(doc.add_paragraph(style="List Number"), item, size=10)
        elif t == "table":
            _table(doc, b)
        elif t == "image":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(io.BytesIO(b["data"]), width=Cm(b.get("width_cm", 16)))
            if b.get("caption"):
                cp = doc.add_paragraph()
                _runs(cp, b["caption"], size=9, color=GREY)
        else:
            raise ValueError(f"unknown block {t}")
    return doc


def to_bytes(blocks):
    buf = io.BytesIO()
    render(blocks).save(buf)
    return buf.getvalue()
