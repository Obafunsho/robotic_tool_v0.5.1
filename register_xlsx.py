"""
Assumptions register as an Excel workbook for clinical review, and reading it back.

Export: one row per procedure and assumption (multi-port and single-port), with a stable
ID, specialty, procedure, assumption, value, unit, source, evidence grade and status,
then blank review columns: Decision (agree / change to / needs evidence), Change to and
Comment. The ID and Value columns are locked (sheet protection, no password) so they
cannot be edited by accident; the review columns are unlocked.

Import: rows marked "change to" with a value become changes, keyed by ID; rows marked
"needs evidence" and any comments are kept for the business case. Changes are applied
to the case configuration by apply_changes() and appear as user-modified assumptions.
"""
import io
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

import defaults as D
import engine as E
import register as RG

COLUMNS = ["ID", "Specialty", "Procedure", "Assumption", "Value", "Unit", "Source", "Evidence grade", "Status",
           "Decision", "Change to", "Comment"]
WIDTHS = [30, 22, 38, 40, 12, 18, 60, 26, 18, 18, 14, 40]
DECISIONS = ["agree", "change to", "needs evidence"]
LOCKED = {"ID", "Value"}
NAVY = "1F3864"
FONT = "Arial"


def register_rows():
    """Every row for the new procedure list: multi-port platform, then single-port, de-duplicated by ID."""
    seen, out = set(), []
    for platform in ("platform_b", "platform_d"):
        for r in RG.detail_rows(E.build_config("large_1", platform, template=None)):
            if r["id"] not in seen:
                seen.add(r["id"])
                out.append(r)
    return out


def export(path=None):
    rows = register_rows()
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    lines = [
        ("Robotic Programme Decision Tool: assumptions register for clinical review", True),
        (f"Engine version {D.ENGINE_VERSION}. Exported {date.today().strftime('%d %B %Y')}.", False),
        ("", False),
        ("How to review", True),
        ("1. Work in the 'Register' sheet. Each row is one assumption for one procedure.", False),
        ("2. In 'Decision', choose agree, change to, or needs evidence.", False),
        ("3. If you choose 'change to', enter the new value in 'Change to' in the same unit as 'Unit'.", False),
        ("   Shares and proportions may be entered as 0.45 or 45%. Whole day: yes or no. Complexity: High, "
         "Medium or Low.", False),
        ("4. Add any comment in 'Comment'.", False),
        ("5. Only the Decision, Change to and Comment columns need editing. ID and Value are locked so they "
         "cannot be changed by accident (the sheet is protected without a password).", False),
        ("6. Save the file and send it back; it is read into the tool and every change is flagged as "
         "user-modified in the business case.", False),
        ("", False),
        ("Example of a completed row (format only, not a recommendation)", True),
        ("ID: colectomy.cases_per_session | Value: 1.5 | Decision: change to | Change to: 1 | "
         "Comment: two colectomies a day on our lists", False),
        ("", False),
        ("Status key", True),
        ("Not yet modelled: 30-day benefit set to zero until the evidence is reviewed. Zero is a placeholder, "
         "not a finding that the robot has no benefit.", False),
        ("Flagged for review: carried over from an earlier procedure group and needs specialty review.", False),
        ("Placeholder: no adequate source yet. Context only: national data for context, not the value itself.",
         False),
        ("Sourced / Partly sourced: from a cited publication; Assumption: a modelling choice; Clinical lead "
         "input: from Aneel's procedure list.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = readme.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=12 if i == 1 else 10, color=NAVY if bold else "000000")
    readme.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Register")
    for j, (name, w) in enumerate(zip(COLUMNS, WIDTHS), start=1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = w
    review_fill = PatternFill("solid", fgColor="FFF8DC")
    status_fill = {"Not yet modelled": "FCE4D6", "Flagged for review": "FFF2CC", "Sourced": "E2EFDA",
                   "Partly sourced": "E2EFDA"}
    for i, r in enumerate(rows, start=2):
        values = [r["id"], r["specialty"], r["procedure"], r["assumption"], r["value"], r["unit"], r["source"],
                  r["grade"], r["status"], None, None, None]
        for j, (name, v) in enumerate(zip(COLUMNS, values), start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.protection = Protection(locked=name in LOCKED)
            if name in ("Decision", "Change to", "Comment"):
                c.fill = review_fill
            if name == "Status" and v in status_fill:
                c.fill = PatternFill("solid", fgColor=status_fill[v])
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "Choose agree, change to or needs evidence", "Decision"
    ws.add_data_validation(dv)
    dv.add(f"J2:J{len(rows) + 1}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:L{len(rows) + 1}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False      # filtering stays allowed
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    wb.active = 1
    if path:
        wb.save(path)
        return path
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Reading a reviewed register back
# ---------------------------------------------------------------------------
def _parse_value(id_, text):
    field = id_.split(".")[-1]
    s = str(text).strip()
    if field == "whole_day":
        if s.lower() in ("yes", "y", "true", "1"):
            return True
        if s.lower() in ("no", "n", "false", "0"):
            return False
        raise ValueError("whole day must be yes or no")
    if field == "complexity":
        v = s.capitalize()
        if v not in D.COMPLEXITY_CASES:
            raise ValueError("complexity must be High, Medium or Low")
        return v
    pct = s.endswith("%")
    v = float(s.rstrip("%").replace(",", "").replace("£", "").strip())
    return v / 100 if pct else v


def load(data):
    """Read a reviewed register (path or bytes). Returns dict(changes, needs_evidence, comments, errors)."""
    wb = load_workbook(io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else data, data_only=True)
    if "Register" not in wb.sheetnames:
        return dict(changes={}, needs_evidence=[], comments={}, errors=["No 'Register' sheet found."])
    ws = wb["Register"]
    head = [c.value for c in ws[1]]
    missing = [c for c in ("ID", "Decision", "Change to", "Comment") if c not in head]
    if missing:
        return dict(changes={}, needs_evidence=[], comments={}, errors=[f"Missing columns: {', '.join(missing)}."])
    ix = {name: head.index(name) for name in head if name}
    known = {r["id"]: r for r in register_rows()}
    changes, needs, comments, errors = {}, [], {}, []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        id_ = row[ix["ID"]]
        if not id_:
            continue
        dec = (row[ix["Decision"]] or "").strip().lower()
        new = row[ix["Change to"]]
        com = row[ix["Comment"]]
        if com:
            comments[id_] = str(com)
        if id_ not in known:
            if dec:
                errors.append(f"Row {n}: unknown ID '{id_}' ignored.")
            continue
        if dec == "change to":
            if new in (None, ""):
                errors.append(f"Row {n} ({id_}): 'change to' without a value; ignored.")
                continue
            try:
                changes[id_] = _parse_value(id_, new)
            except ValueError as exc:
                errors.append(f"Row {n} ({id_}): {exc}; ignored.")
        elif dec == "needs evidence":
            needs.append(id_)
        elif dec and dec != "agree":
            errors.append(f"Row {n} ({id_}): decision '{dec}' not recognised; ignored.")
    return dict(changes=changes, needs_evidence=needs, comments=comments, errors=errors)


def apply_changes(overrides, changes, preset):
    """Apply register changes to build_config overrides. Demand changes apply to the
    preset's hospital size; whole-day or complexity changes reset cases per session to
    the default for the new setting unless cases per session is also changed."""
    size = "medium" if preset.startswith("medium") else "large"
    groups = overrides.setdefault("groups", {})
    for id_, v in changes.items():
        parts = id_.split(".")
        if parts[0] == "unit_costs":
            overrides.setdefault("unit_costs", {})[parts[1]] = v
        elif parts[0] == "programme":
            overrides.setdefault("programme", {})[parts[1]] = v
        elif len(parts) == 2 and parts[1].startswith("demand_"):
            g = groups.setdefault(parts[0], {})
            g[f"cases_{parts[1][7:]}"] = v
            if parts[1] == f"demand_{size}":
                g["volume"] = v
        elif len(parts) == 2:
            groups.setdefault(parts[0], {})[parts[1]] = v
        elif len(parts) == 3 and parts[1] == "share":
            g = groups.setdefault(parts[0], {})
            base = dict(D.SP_GROUPS.get(parts[0], {}).get("comparators", {}))
            comps = dict(g.get("comparators", base))
            comps[parts[2]] = v
            g["comparators"] = comps
        elif len(parts) == 3:
            g = groups.setdefault(parts[0], {})
            ref = D.SP_GROUPS.get(parts[0], {})
            if parts[2] == "extra_minutes":
                em = dict(g.get("extra_minutes", ref.get("extra_minutes", {})))
                em[parts[1]] = v
                g["extra_minutes"] = em
            else:
                import copy
                red = copy.deepcopy(g.get("reductions", ref.get("reductions", {})))
                red.setdefault(parts[1], {})[parts[2]] = v
                g["reductions"] = red
    for gk, g in groups.items():
        if ("whole_day" in g or "complexity" in g) and "cases_per_session" not in g:
            ref = D.PROCEDURE_GROUPS.get(gk) or D.SP_GROUPS.get(gk)
            if ref:
                g["cases_per_session"] = D.default_cases_per_session(g.get("complexity", ref["complexity"]),
                                                                     g.get("whole_day", ref["whole_day"]))
    return overrides
